import os
import sys
from pathlib import Path
import openpyxl

# Add project root to sys.path if not present
BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from backend.database import get_db_connection, init_db

# Configure stdout for UTF-8 on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"


def seed_events():
    """Seeds the events table from data/events_seed.xlsx if the table is empty."""
    # Ensure database schema exists
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events already exist
    cursor.execute("SELECT COUNT(*) FROM events")
    count = cursor.fetchone()[0]
    if count > 0:
        print(f"Events table already has {count} rows. Skipping seed.")
        conn.close()
        return

    # Check for seed file existence
    if not SEED_FILE_PATH.exists():
        raise FileNotFoundError(
            f"Seed file not found at {SEED_FILE_PATH}. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    # Load workbook and sheet
    wb = openpyxl.load_workbook(SEED_FILE_PATH, data_only=True)
    if "Events" not in wb.sheetnames:
        raise ValueError("Sheet 'Events' not found in seed workbook.")

    ws = wb["Events"]

    # Read rows (row 1 is header: id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise ValueError("No data rows found in seed workbook.")

    header = [str(col).strip().lower() if col is not None else "" for col in rows[0]]
    expected_header = ["id", "title", "category", "description", "event_date", "event_time", "venue", "capacity", "fee", "icon"]

    # Validate header row
    for exp in expected_header:
        if exp not in header:
            raise ValueError(f"Missing expected header column: {exp}")

    col_idx = {name: header.index(name) for name in expected_header}

    inserted_count = 0
    for row in rows[1:]:
        if not row or all(v is None for v in row):
            continue

        event_id = row[col_idx["id"]]
        title = row[col_idx["title"]]
        category = row[col_idx["category"]]
        description = row[col_idx["description"]]
        event_date = str(row[col_idx["event_date"]])[:10] if row[col_idx["event_date"]] is not None else ""
        event_time = str(row[col_idx["event_time"]]) if row[col_idx["event_time"]] is not None else ""
        venue = row[col_idx["venue"]]
        capacity = int(row[col_idx["capacity"]]) if row[col_idx["capacity"]] is not None else 0
        fee = int(row[col_idx["fee"]]) if row[col_idx["fee"]] is not None else 0
        icon = row[col_idx["icon"]]

        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon),
        )
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {SEED_FILE_PATH}.")


def display_events():
    """Fetches and displays all events in the events table."""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    rows = cursor.fetchall()
    conn.close()

    print(f"\n--- Events in fest.db ({len(rows)} rows) ---")
    header_fmt = "{:<4} | {:<24} | {:<12} | {:<12} | {:<10} | {:<20} | {:<8} | {:<6} | {:<4}"
    print(header_fmt.format("ID", "Title", "Category", "Date", "Time", "Venue", "Capacity", "Fee", "Icon"))
    print("-" * 115)
    for r in rows:
        print(header_fmt.format(
            r["id"],
            r["title"][:24],
            r["category"][:12],
            r["event_date"][:12],
            r["event_time"][:10],
            r["venue"][:20],
            r["capacity"],
            r["fee"],
            r["icon"] or ""
        ))
    print("-" * 115)


if __name__ == "__main__":
    seed_events()
    display_events()
